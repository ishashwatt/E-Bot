import os
import re
from typing import List, Dict, Any, Tuple
import pypdf
import pdfplumber
import openpyxl

class AttachmentParser:
    """
    Parses PDF, Excel, and Text attachments to extract text and identify exact page/row hits for search criteria.
    """

    @staticmethod
    def parse_pdf(file_path: str, search_terms: List[str]) -> Dict[str, Any]:
        result = {
            "text": "",
            "total_pages": 0,
            "matches": [],
            "preview": ""
        }
        
        if not os.path.exists(file_path):
            return result

        try:
            with pdfplumber.open(file_path) as pdf:
                result["total_pages"] = len(pdf.pages)
                all_text = []

                for page_idx, page in enumerate(pdf.pages, start=1):
                    page_text = page.extract_text() or ""
                    all_text.append(f"--- Page {page_idx} ---\n" + page_text)

                    for term in search_terms:
                        if not term:
                            continue
                        pattern = re.compile(re.escape(term), re.IGNORECASE)
                        for line_num, line in enumerate(page_text.splitlines(), start=1):
                            if pattern.search(line):
                                result["matches"].append({
                                    "term": term,
                                    "page": page_idx,
                                    "line": line_num,
                                    "snippet": line.strip()
                                })

                full_text = "\n".join(all_text)
                result["text"] = full_text
                result["preview"] = full_text[:1000]

        except Exception as e:
            try:
                reader = pypdf.PdfReader(file_path)
                result["total_pages"] = len(reader.pages)
                all_text = []
                for page_idx, page in enumerate(reader.pages, start=1):
                    page_text = page.extract_text() or ""
                    all_text.append(f"--- Page {page_idx} ---\n" + page_text)
                    for term in search_terms:
                        if term and re.search(re.escape(term), page_text, re.IGNORECASE):
                            result["matches"].append({
                                "term": term,
                                "page": page_idx,
                                "line": 1,
                                "snippet": f"Found '{term}' on Page {page_idx}"
                            })
                full_text = "\n".join(all_text)
                result["text"] = full_text
                result["preview"] = full_text[:1000]
            except Exception as fallback_err:
                result["text"] = f"Error reading PDF: {str(e)} / {str(fallback_err)}"

        return result

    @staticmethod
    def parse_excel(file_path: str, search_terms: List[str]) -> Dict[str, Any]:
        result = {
            "text": "",
            "total_pages": 1,
            "matches": [],
            "preview": ""
        }
        
        if not os.path.exists(file_path):
            return result

        try:
            wb = openpyxl.load_workbook(file_path, data_only=True)
            text_lines = []
            
            for sheetname in wb.sheetnames:
                sheet = wb[sheetname]
                text_lines.append(f"--- Sheet: {sheetname} ---")
                
                for row_idx, row in enumerate(sheet.iter_rows(values_only=True), start=1):
                    row_str = " | ".join([str(cell) for cell in row if cell is not None])
                    if not row_str.strip():
                        continue
                    text_lines.append(row_str)

                    for term in search_terms:
                        if term and re.search(re.escape(term), row_str, re.IGNORECASE):
                            result["matches"].append({
                                "term": term,
                                "sheet": sheetname,
                                "row": row_idx,
                                "snippet": row_str[:200]
                            })

            full_text = "\n".join(text_lines)
            result["text"] = full_text
            result["preview"] = full_text[:1000]
        except Exception as e:
            result["text"] = f"Error reading Excel file: {str(e)}"

        return result

    @staticmethod
    def parse_csv(file_path: str, search_terms: List[str]) -> Dict[str, Any]:
        import csv
        result = {
            "text": "",
            "total_pages": 1,
            "matches": [],
            "preview": ""
        }
        if not os.path.exists(file_path):
            return result

        try:
            text_lines = []
            with open(file_path, mode="r", encoding="utf-8", errors="ignore") as f:
                reader = csv.reader(f)
                for row_idx, row in enumerate(reader, start=1):
                    row_str = " | ".join([str(c).strip() for c in row if c])
                    if not row_str:
                        continue
                    text_lines.append(row_str)
                    for term in search_terms:
                        if term and re.search(re.escape(term), row_str, re.IGNORECASE):
                            result["matches"].append({
                                "term": term,
                                "sheet": "CSV",
                                "row": row_idx,
                                "snippet": row_str[:200]
                            })
            full_text = "\n".join(text_lines)
            result["text"] = full_text
            result["preview"] = full_text[:1000]
        except Exception as e:
            result["text"] = f"Error reading CSV: {str(e)}"
        return result

    @classmethod
    def parse_attachment(cls, file_path: str, search_terms: List[str]) -> Dict[str, Any]:
        ext = os.path.splitext(file_path)[1].lower()
        if ext == ".pdf":
            return cls.parse_pdf(file_path, search_terms)
        elif ext in [".xlsx", ".xlsm"]:
            return cls.parse_excel(file_path, search_terms)
        elif ext in [".csv", ".tsv", ".txt"]:
            return cls.parse_csv(file_path, search_terms)
        else:
            try:
                with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                    content = f.read()
                    matches = []
                    for term in search_terms:
                        if term and re.search(re.escape(term), content, re.IGNORECASE):
                            matches.append({"term": term, "snippet": f"Found in text: {term}"})
                    return {
                        "text": content,
                        "total_pages": 1,
                        "matches": matches,
                        "preview": content[:1000]
                    }
            except Exception as e:
                return {"text": f"Error: {e}", "total_pages": 0, "matches": [], "preview": ""}

